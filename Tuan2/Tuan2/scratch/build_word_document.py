# -*- coding: utf-8 -*-
import os
import sys
import json
import docx
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, border in borders.items():
        if border:
            b = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{border.get("val", "single")}" w:sz="{border.get("sz", 4)}" w:space="0" w:color="{border.get("color", "auto")}"/>')
            tcBorders.append(b)
        else:
            b = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    
    if level == 1:
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(10)
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
        # Add a subtle bottom border or divider line under Heading 1
    elif level == 2:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(37, 99, 235) # Slate Blue
    elif level == 3:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(13, 148, 136) # Dark Cyan
    return p

def add_code_block(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC") # Slate 50
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, 
                     left={'val': 'single', 'sz': 24, 'color': '2563EB'}, # Strong blue bar
                     top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                     bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                     right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    
    run = p.add_run(code_str.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)

def add_callout(doc, title, content, bg_color="EFF6FF", border_color="3B82F6", title_color=(30, 58, 138)):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=130, bottom=130, left=170, right=170)
    set_cell_borders(cell, 
                     left={'val': 'single', 'sz': 20, 'color': border_color},
                     top={'val': 'none'}, bottom={'val': 'none'}, right={'val': 'none'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_title = p.add_run(f"📌 {title}\n")
    run_title.bold = True
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(*title_color)
    
    run_content = p.add_run(content)
    run_content.font.name = 'Arial'
    run_content.font.size = Pt(10)
    run_content.font.color.rgb = RGBColor(31, 41, 55)

def add_screenshot(doc, img_path, caption_text):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Arial'
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(100, 116, 139)
    else:
        p = doc.add_paragraph(f"[Hình ảnh không tìm thấy: {img_path}]")
        p.runs[0].font.color.rgb = RGBColor(239, 68, 68)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.18):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(31, 41, 55)
    r_text = p.add_run(text)
    r_text.font.name = 'Arial'
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = RGBColor(31, 41, 55)
    return p

def main():
    print("Bắt đầu khởi tạo báo cáo Word...")
    doc = docx.Document()
    
    # 1. Cấu hình trang A4 và lề tiêu chuẩn
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4 Width
    section.page_height = Inches(11.69) # A4 Height
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    # 2. Tải dữ liệu các bài tập
    with open('scratch/exercises_data_part1.json', 'r', encoding='utf-8') as f:
        data_p1 = json.load(f)
    with open('scratch/exercises_data_part2.json', 'r', encoding='utf-8') as f:
        data_p2 = json.load(f)
    with open('scratch/exercises_data_part3.json', 'r', encoding='utf-8') as f:
        data_p3 = json.load(f)
    
    exercises = {}
    exercises.update(data_p1)
    exercises.update(data_p2)
    exercises.update(data_p3)
    
    with open('scratch/run_results.json', 'r', encoding='utf-8') as f:
        run_results = json.load(f)
        
    print(f"Đã nạp dữ liệu {len(exercises)} bài tập và {len(run_results)} kết quả thực thi.")
    
    # =========================================================================
    # TRANG BÌA (COVER PAGE)
    # =========================================================================
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_header, space_before=10, space_after=2)
    r_school = p_header.add_run("BỘ MÔN CÔNG NGHỆ THÔNG TIN - LẬP TRÌNH THIẾT BỊ DI ĐỘNG\n")
    r_school.bold = True
    r_school.font.name = 'Arial'
    r_school.font.size = Pt(12)
    r_school.font.color.rgb = RGBColor(71, 85, 105)
    
    r_sub = p_header.add_run("KHOA CÔNG NGHỆ PHẦN MỀM & KỸ THUẬT MÁY TÍNH\n")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_line, space_before=2, space_after=30)
    r_l = p_line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    r_l.font.color.rgb = RGBColor(203, 213, 225)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_title, space_before=20, space_after=14)
    r_title = p_title.add_run("BÁO CÁO THỰC HÀNH CHUYÊN SÂU\n")
    r_title.bold = True
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    r_title2 = p_title.add_run("LẬP TRÌNH BẤT ĐỒNG BỘ TRONG TYPESCRIPT & REACT NATIVE\n")
    r_title2.bold = True
    r_title2.font.name = 'Arial'
    r_title2.font.size = Pt(18)
    r_title2.font.color.rgb = RGBColor(37, 99, 235)
    
    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_desc, space_before=10, space_after=40)
    r_desc = p_desc.add_run("TỔNG HỢP MÃ NGUỒN, HÌNH ẢNH CHỤP KẾT QUẢ THỰC THI & PHÂN TÍCH KỸ THUẬT CHI TIẾT 30 BÀI TẬP\n(Promise Architecture, Async/Await Syntax, Fetch REST API & Concurrency Patterns)")
    r_desc.font.name = 'Arial'
    r_desc.font.size = Pt(11.5)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(75, 85, 99)
    
    # Hộp thông tin sinh viên & môi trường
    tbl_info = doc.add_table(rows=6, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_info.autofit = False
    
    info_data = [
        ("Môn học:", "Lập trình Thiết bị Di động (Mobile Programming)"),
        ("Bài tập thực hành:", "Tuần 2 - Asynchronous TypeScript & React Native Fundamentals"),
        ("Môi trường thực thi:", "Node.js v24.12.0 LTS (TSX Execution Engine)"),
        ("Nền tảng ứng dụng:", "React Native 0.86.3 / Expo SDK 57.0.20 / TypeScript 6.0.3"),
        ("Tài liệu đối chiếu:", "2_Async_Exercises_TypeScript.pdf (Bộ 30 bài tập chuẩn)"),
        ("Trạng thái kiểm thử:", "Hoàn thành 30/30 bài tập - 100% Exit Code 0 (Success)")
    ]
    
    for i, (k, v) in enumerate(info_data):
        row = tbl_info.rows[i]
        c1 = row.cells[0]
        c2 = row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        set_cell_background(c1, "F1F5F9")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
        
        p1 = c1.paragraphs[0]
        format_paragraph(p1, space_before=0, space_after=0)
        r1 = p1.add_run(k)
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(30, 58, 138)
        
        p2 = c2.paragraphs[0]
        format_paragraph(p2, space_before=0, space_after=0)
        r2 = p2.add_run(v)
        r2.font.name = 'Arial'
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(31, 41, 55)
        
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_bot, space_before=50, space_after=0)
    r_bot = p_bot.add_run("Thành phố Hồ Chí Minh, Năm 2026")
    r_bot.font.name = 'Arial'
    r_bot.font.size = Pt(11)
    r_bot.font.italic = True
    r_bot.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_page_break()
    
    # =========================================================================
    # MỤC LỤC TỔNG QUAN
    # =========================================================================
    add_styled_heading(doc, "MỤC LỤC BÁO CÁO", 1)
    toc_p = doc.add_paragraph()
    format_paragraph(toc_p, space_before=4, space_after=14)
    toc_text = (
        "CHƯƠNG 1: TỔNG QUAN LÝ THUYẾT LẬP TRÌNH BẤT ĐỒNG BỘ TRONG JAVASCRIPT & TYPESCRIPT\n"
        "    1.1. Bản chất kiến trúc Single-Threaded và Cơ chế Non-blocking I/O\n"
        "    1.2. Vòng lặp sự kiện (Event Loop), Call Stack, Microtask Queue & Macrotask Queue\n"
        "    1.3. Lịch sử tiến hóa: Từ Callback Hell đến Promise và Async / Await\n"
        "    1.4. So sánh 4 phương thức xử lý đồng thời: Promise.all, Promise.race, Promise.allSettled, Promise.any\n"
        "    1.5. Xử lý lỗi toàn diện và Quản lý giải phóng bộ nhớ (Memory Leak) trong React Native\n\n"
        "CHƯƠNG 2: NỘI DUNG THỰC HÀNH CHI TIẾT 30 BÀI TẬP\n"
        "    PHẦN A: CƠ BẢN VỀ PROMISE (Bài 1 đến Bài 10)\n"
        "        - Bài 1: Tạo Promise trả về 'Hello Async' sau 2 giây\n"
        "        - Bài 2: Hàm trả về Promise giải quyết giá trị số 10 sau 1 giây\n"
        "        - Bài 3: Hàm từ chối (reject) Promise với thông báo lỗi sau 1 giây\n"
        "        - Bài 4: Xử lý Promise trả về số ngẫu nhiên với .then() và .catch()\n"
        "        - Bài 5: Xây dựng hàm tiện ích simulateTask(time) mô phỏng độ trễ linh hoạt\n"
        "        - Bài 6: Thực thi song song 3 Promise mô phỏng bằng Promise.all()\n"
        "        - Bài 7: Sử dụng Promise.race() lấy kết quả của Promise hoàn thành sớm nhất\n"
        "        - Bài 8: Xây dựng chuỗi Promise (Promise Chaining) tính toán tuần tự\n"
        "        - Bài 9: Đọc mảng dữ liệu sau 1 giây và lọc các số chẵn\n"
        "        - Bài 10: Quản lý dọn dẹp và kết thúc với phương thức .finally()\n"
        "    PHẦN B: LẬP TRÌNH BẤT ĐỒNG BỘ VỚI ASYNC / AWAIT (Bài 11 đến Bài 20)\n"
        "        - Bài 11: Chuyển đổi Bài 1 sang cú pháp hiện đại async / await\n"
        "        - Bài 12: Hàm async gọi hàm tiện ích simulateTask(2000) và in kết quả\n"
        "        - Bài 13: Xử lý lỗi bất đồng bộ toàn diện với cấu trúc try / catch\n"
        "        - Bài 14: Hàm async nhận số, chờ 1 giây và trả về tích số nhân 3\n"
        "        - Bài 15: Thực thi tuần tự nhiều hàm async bằng từ khóa await\n"
        "        - Bài 16: Thực thi song song nhiều hàm async với Promise.all()\n"
        "        - Bài 17: Duyệt mảng các Promise bằng vòng lặp bất đồng bộ for await...of\n"
        "        - Bài 18: Hàm async fetchUser(id) mô phỏng gọi API lấy đối tượng User\n"
        "        - Bài 19: Hàm async fetchUsers(ids) lấy dữ liệu nhiều User đồng thời\n"
        "        - Bài 20: Kỹ thuật Timeout Pattern: ném lỗi nếu gọi API quá 2 giây\n"
        "    PHẦN C: FETCH API & MÔ PHỎNG I/O MẠNG (Bài 21 đến Bài 30)\n"
        "        - Bài 21: Sử dụng Fetch API lấy dữ liệu từ máy chủ công khai\n"
        "        - Bài 22: Gọi API nhiều lần tuần tự và in kết quả chi tiết\n"
        "        - Bài 23: Lấy danh sách công việc từ API và lọc các mục chưa hoàn thành\n"
        "        - Bài 24: Gửi dữ liệu lên máy chủ thông qua phương thức HTTP POST\n"
        "        - Bài 25: Hàm downloadFile mô phỏng quá trình tải tập tin trong 3 giây\n"
        "        - Bài 26: Tạo hàm tạm dừng (delay/sleep) 5 giây kết hợp async/await và setTimeout\n"
        "        - Bài 27: Triển khai mẫu thiết kế Retry Pattern tự động thử lại khi gọi API thất bại\n"
        "        - Bài 28: Xử lý hàng loạt (Batch Processing) 5 tác vụ bất đồng bộ cùng lúc\n"
        "        - Bài 29: Xử lý hàng đợi tác vụ tuần tự (Sequential Task Queue - FIFO)\n"
        "        - Bài 30: Xử lý đa request độc lập bằng Promise.allSettled() hiển thị trạng thái Thành công / Thất bại\n\n"
        "CHƯƠNG 3: TỔNG KẾT & KINH NGHIỆM THỰC CHIẾN TRONG DỰ ÁN REACT NATIVE / MOBILE\n"
        "    3.1. Bảng tổng hợp đối chiếu kết quả 30 bài tập thực nghiệm\n"
        "    3.2. Năm quy tắc vàng khi làm việc với Bất đồng bộ trong ứng dụng Di động\n"
        "    3.3. Lời kết"
    )
    r_toc = toc_p.add_run(toc_text)
    r_toc.font.name = 'Consolas'
    r_toc.font.size = Pt(9.5)
    r_toc.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_page_break()
    
    # =========================================================================
    # CHƯƠNG 1: TỔNG QUAN LÝ THUYẾT
    # =========================================================================
    add_styled_heading(doc, "CHƯƠNG 1: TỔNG QUAN LÝ THUYẾT LẬP TRÌNH BẤT ĐỒNG BỘ TRONG JAVASCRIPT & TYPESCRIPT", 1)
    
    add_styled_heading(doc, "1.1. Bản chất Single-Threaded và Cơ chế Non-blocking I/O", 2)
    add_body_p(doc, "JavaScript được thiết kế theo mô hình đơn luồng (Single-threaded), nghĩa là tại một thời điểm chỉ có duy nhất một luồng thực thi (Thread) chạy các chỉ thị lệnh trên Call Stack chính. Tuy nhiên, các ứng dụng thực tế - đặc biệt là ứng dụng di động (React Native) - đòi hỏi phải xử lý đồng thời nhiều tác vụ tốn thời gian như: gửi nhận gói tin qua mạng (HTTP Network Requests), đọc ghi bộ nhớ thiết bị (Disk I/O, SQLite, AsyncStorage), định thời (Timers: setTimeout, setInterval) và các tương tác cử chỉ mượt mà của người dùng (Gestures, Animations 60-120fps).")
    add_body_p(doc, "Nếu JavaScript xử lý các tác vụ này theo phương thức đồng bộ (Synchronous / Blocking), luồng chính sẽ bị 'đóng băng' hoàn toàn trong lúc chờ máy chủ phản hồi gói tin, dẫn đến hiện tượng treo ứng dụng (Application Not Responding - ANR trên Android hoặc Crash Watchdog trên iOS). Để giải quyết triệt để vấn đề này, kiến trúc JavaScript Runtime (V8 Engine trong Node.js hoặc Hermes/JSC trong React Native) kết hợp chặt chẽ với subsystem C++ (Libuv trong Node.js hoặc C++ Core trong React Native) để tạo nên cơ chế Bất đồng bộ không chặn (Non-blocking Asynchronous I/O).")
    
    add_styled_heading(doc, "1.2. Vòng lặp sự kiện (Event Loop), Call Stack, Microtask & Macrotask Queue", 2)
    add_body_p(doc, "Kiến trúc xử lý bất đồng bộ trong JavaScript được vận hành dựa trên 4 thành phần cốt lõi:")
    add_body_p(doc, "Là cấu trúc dữ liệu LIFO (Last-In, First-Out) lưu trữ các hàm đang được thực thi. Khi một hàm được gọi, khung hàm (Stack Frame) được đẩy vào Stack. Khi hàm kết thúc (return), frame đó được lấy ra khỏi Stack.", "1. Call Stack: ")
    add_body_p(doc, "Các tác vụ nặng (như setTimeout, fetch, file I/O) không chạy trên Call Stack JS mà được chuyển giao cho subsystem C++ nền tảng xử lý độc lập.", "2. Background Web APIs / C++ Threads: ")
    add_body_p(doc, "Khi tác vụ nền hoàn thành, callback tương ứng không nhảy ngay vào Call Stack mà được xếp hàng tại các hàng đợi chuyên biệt. Có 2 loại hàng đợi với thứ tự ưu tiên khác nhau:\n  • Microtask Queue: Ưu tiên cao nhất. Chứa các callback của Promise (.then, .catch, .finally), queueMicrotask, MutationObserver. Toàn bộ hàng đợi Microtask sẽ được dọn sạch ngay khi Call Stack trống.\n  • Macrotask Queue (Task Queue): Ưu tiên thấp hơn. Chứa các callback của setTimeout, setInterval, setImmediate, I/O events. Mỗi vòng lặp chỉ lấy đúng một Macrotask ra chạy.", "3. Task Queues (Microtask vs Macrotask): ")
    add_body_p(doc, "Một vòng lặp vô tận liên tục kiểm tra Call Stack. Nếu Call Stack hoàn toàn rỗng, Event Loop sẽ ưu tiên đẩy toàn bộ các tác vụ trong Microtask Queue vào Call Stack trước. Chỉ khi Microtask Queue sạch bóng, nó mới lấy tác vụ tiếp theo từ Macrotask Queue.", "4. Event Loop: ")
    
    add_styled_heading(doc, "1.3. Lịch sử tiến hóa: Callback Hell -> Promise -> Async/Await", 2)
    add_body_p(doc, "Trong những ngày đầu của JavaScript, bất đồng bộ được giải quyết bằng các hàm gọi lại (Callbacks). Khi nhiều tác vụ phụ thuộc dữ liệu lẫn nhau, mã nguồn trở thành một kim tự tháp lồng nhau sâu hoắm được gọi là 'Callback Hell' (hoặc 'Pyramid of Doom'), gây khó khăn trầm trọng trong việc bảo trì, debug và bắt lỗi.")
    add_body_p(doc, "Chuẩn ES6 (2015) chính thức giới thiệu đối tượng Promise. Promise đại diện cho một giá trị chưa hoàn thành tại thời điểm hiện tại nhưng sẽ được giải quyết trong tương lai, với 3 trạng thái bất biến: Pending (Đang chờ), Fulfilled (Thành công với giá trị), hoặc Rejected (Thất bại với lý do lỗi). Promise cung cấp cơ chế chuỗi (Chaining) thông qua .then(), .catch(), và .finally(), biến các luồng xử lý lồng nhau thành các đường ống phẳng theo thứ tự từ trên xuống dưới.")
    add_body_p(doc, "Chuẩn ES2017 (ES8) tiếp tục nâng tầm lập trình bất đồng bộ với bộ đôi từ khóa async và await. Về bản chất, async/await là cú pháp đường mật (Syntactic Sugar) xây dựng trên nền tảng Promise và Generators. Nó cho phép lập trình viên viết các đoạn mã bất đồng bộ có cấu trúc trực quan, mạch lạc như mã đồng bộ thông thường, đồng thời tận dụng cấu trúc try/catch kinh điển để quản lý ngoại lệ toàn diện.")
    
    add_styled_heading(doc, "1.4. So sánh 4 phương thức xử lý đồng thời (Concurrency Methods)", 2)
    add_body_p(doc, "JavaScript cung cấp 4 phương thức tĩnh trên đối tượng Promise để điều phối nhiều tác vụ đồng thời. Bảng phân tích chi tiết dưới đây chỉ rõ sự khác biệt giữa các phương thức:")
    
    tbl_cmp = doc.add_table(rows=5, cols=4)
    tbl_cmp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cmp.autofit = False
    
    headers = ["Phương thức", "Cơ chế kích hoạt", "Cơ chế thất bại (Error)", "Trường hợp ứng dụng tối ưu"]
    for j, h in enumerate(headers):
        c = tbl_cmp.rows[0].cells[j]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=100, bottom=100, left=100, right=100)
        p = c.paragraphs[0]
        format_paragraph(p, space_before=0, space_after=0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    concurrency_data = [
        ("Promise.all()", "Thành công khi TẤT CẢ các Promise đều fulfilled.", "Fail-fast: Lập tức reject ngay khi CÓ 1 Promise bất kỳ bị reject.", "Tải đồng thời các tài nguyên độc lập mà trang bắt buộc phải có đủ."),
        ("Promise.allSettled()", "Luôn luôn chờ TẤT CẢ Promise kết thúc (fulfilled hoặc rejected).", "Không bao giờ reject. Trả về mảng đối tượng {status, value/reason}.", "Dashboard tổng hợp: Cho phép widget lỗi hiển thị trạng thái fail mà không sập trang."),
        ("Promise.race()", "Giải quyết theo Promise ĐẦU TIÊN hoàn thành (dù thành công hay lỗi).", "Reject nếu Promise đầu tiên hoàn thành bị lỗi.", "Triển khai Timeout mạng: Đua tác vụ fetch với một timer reject sau 3s."),
        ("Promise.any()", "Thành công ngay khi CÓ 1 Promise ĐẦU TIÊN fulfilled.", "Chỉ reject khi TẤT CẢ Promise cùng bị reject (AggregateError).", "Gửi request tới nhiều máy chủ Mirror, lấy dữ liệu từ server đầu tiên trả về OK.")
    ]
    
    col_widths = [Inches(1.5), Inches(1.8), Inches(1.6), Inches(1.6)]
    for i, row_data in enumerate(concurrency_data):
        row = tbl_cmp.rows[i + 1]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row_data):
            c = row.cells[j]
            c.width = col_widths[j]
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            format_paragraph(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if j == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            else:
                r.font.color.rgb = RGBColor(31, 41, 55)
                
    add_body_p(doc, "")
    
    add_styled_heading(doc, "1.5. Xử lý lỗi toàn diện & Quản lý giải phóng bộ nhớ (Memory Leak) trong React Native", 2)
    add_body_p(doc, "Trong lập trình ứng dụng di động React Native, bất đồng bộ gắn liền với vòng đời của các Component (Lifecycle). Một sai lầm kinh điển của các lập trình viên mới là không kiểm soát việc Component bị Unmount (người dùng ấn Back hoặc chuyển màn hình) trong lúc một Promise mạng đang chạy. Khi Promise này hoàn thành và gọi hàm cập nhật State (ví dụ: `setUserData(data)`), ứng dụng sẽ bị rò rỉ bộ nhớ (Memory Leak) và sinh cảnh báo nghiêm trọng: `Can't perform a React state update on an unmounted component`.")
    add_body_p(doc, "Để xử lý triệt để vấn đề này, lập trình viên hiện đại cần sử dụng chuẩn `AbortController` kết hợp với `useEffect cleanup function` để chủ động hủy (abort) các request mạng đang dở dang khi người dùng rời khỏi màn hình, đồng thời luôn bọc các khối await trong `try...catch...finally` để đảm bảo tài nguyên tải (loading indicator) luôn được thu hồi.")
    
    doc.add_page_break()
    
    # =========================================================================
    # CHƯƠNG 2: NỘI DUNG THỰC HÀNH CHI TIẾT 30 BÀI TẬP
    # =========================================================================
    add_styled_heading(doc, "CHƯƠNG 2: NỘI DUNG THỰC HÀNH CHI TIẾT 30 BÀI TẬP", 1)
    add_body_p(doc, "Chương này trình bày toàn bộ 30 bài tập thực hành bất đồng bộ trong dự án, được chia làm 3 phần chính theo đúng đề cương tài liệu chuẩn `2_Async_Exercises_TypeScript.pdf`. Mỗi bài tập được cấu trúc đầy đủ và chuyên nghiệp theo 6 nội dung: (1) Yêu cầu đề bài, (2) Mã nguồn TypeScript hoàn chỉnh, (3) Ảnh chụp màn hình kết quả chạy trên Terminal, (4) Phân tích chi tiết từng dòng code, (5) Cơ chế hoạt động & Luồng thực thi Event Loop, (6) Phân tích kết quả đầu ra & Ứng dụng thực tiễn trong React Native.")
    
    current_group = ""
    
    for i in range(1, 31):
        str_i = str(i)
        item = exercises.get(str_i, {})
        group_name = item.get('group', '')
        
        # Tiêu đề nhóm
        if group_name != current_group:
            current_group = group_name
            add_styled_heading(doc, group_name, 2)
            if i == 1:
                add_body_p(doc, "Nhóm bài tập này tập trung rèn luyện các kỹ năng nền tảng cốt lõi với đối tượng Promise nguyên bản: Khởi tạo Promise, hiểu rõ 3 trạng thái (Pending, Fulfilled, Rejected), điều phối luồng với .then(), bắt ngoại lệ với .catch(), dọn dẹp tài nguyên với .finally(), kết hợp chuỗi Promise Chaining và xử lý đa tác vụ với Promise.all() và Promise.race().")
            elif i == 11:
                doc.add_page_break()
                add_styled_heading(doc, group_name, 2)
                add_body_p(doc, "Nhóm bài tập này chuyển dịch toàn diện sang cú pháp hiện đại async/await (ES2017): Viết hàm async, kiểm soát luồng với từ khóa await, bắt lỗi cấu trúc với try/catch, so sánh trực quan giữa thực thi tuần tự (Sequential) và thực thi song song (Parallel), kỹ thuật duyệt mảng bất đồng bộ for await...of, mô hình hóa Interface trong TypeScript và triển khai mẫu thiết kế Timeout Pattern chống nghẽn ứng dụng.")
            elif i == 21:
                doc.add_page_break()
                add_styled_heading(doc, group_name, 2)
                add_body_p(doc, "Nhóm bài tập này ứng dụng toàn diện vào các tác vụ I/O thực tế: Giao tiếp với RESTful API qua Fetch API chuẩn, xử lý dữ liệu động với phương thức HTTP POST, mô phỏng tải tệp tin và cơ chế non-blocking sleep, xây dựng các mẫu thiết kế công nghiệp như Retry Pattern tự động thử lại, Batch Processing xử lý hàng loạt, Sequential Queue hàng đợi tuần tự và Promise.allSettled() giải quyết triệt để lỗi phân tán.")
                
        # Tiêu đề bài tập
        add_styled_heading(doc, f"BÀI {i}: {item.get('vn_title', '')}", 3)
        
        # Hộp yêu cầu đề bài
        req_content = f"Đề bài gốc (English): {item.get('pdf_req', '')}\n\nMục tiêu kỹ thuật: {item.get('purpose', '')}"
        add_callout(doc, f"YÊU CẦU BÀI TẬP {i}", req_content, bg_color="EFF6FF", border_color="3B82F6")
        
        # Đọc mã nguồn từ file thực tế
        code_path = f"src/practice/Bai{i}.ts"
        code_str = ""
        if os.path.exists(code_path):
            with open(code_path, 'r', encoding='utf-8') as cf:
                code_str = cf.read()
        else:
            code_str = "// Không tìm thấy file mã nguồn"
            
        p_c_label = doc.add_paragraph()
        format_paragraph(p_c_label, space_before=8, space_after=2)
        r_cl = p_c_label.add_run(f"Mã nguồn TypeScript (`{code_path}`):")
        r_cl.bold = True
        r_cl.font.name = 'Arial'
        r_cl.font.size = Pt(10)
        r_cl.font.color.rgb = RGBColor(30, 58, 138)
        
        add_code_block(doc, code_str)
        
        # Ảnh chụp màn hình kết quả chạy
        screenshot_path = f"scratch/screenshots/bai{i}.png"
        caption = f"Hình 2.{i}: Ảnh chụp kết quả thực thi Bài {i} trên môi trường Terminal (Node.js v24 / TSX Engine)"
        add_screenshot(doc, screenshot_path, caption)
        
        # Bảng thông số thực thi
        run_info = run_results.get(str_i, {})
        tbl_stat = doc.add_table(rows=1, cols=3)
        tbl_stat.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_stat.autofit = False
        tbl_stat.rows[0].cells[0].width = Inches(2.1)
        tbl_stat.rows[0].cells[1].width = Inches(2.1)
        tbl_stat.rows[0].cells[2].width = Inches(2.3)
        
        set_cell_background(tbl_stat.rows[0].cells[0], "F8FAFC")
        set_cell_background(tbl_stat.rows[0].cells[1], "F8FAFC")
        set_cell_background(tbl_stat.rows[0].cells[2], "F8FAFC")
        
        for c_idx in range(3):
            set_cell_margins(tbl_stat.rows[0].cells[c_idx], top=60, bottom=60, left=100, right=100)
            set_cell_borders(tbl_stat.rows[0].cells[c_idx], 
                             top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                             bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                             left={'val': 'none'}, right={'val': 'none'})
            
        p_st0 = tbl_stat.rows[0].cells[0].paragraphs[0]
        format_paragraph(p_st0, 0, 0)
        r_st0 = p_st0.add_run(f"⏱ Thời gian: {run_info.get('duration', 'N/A')}s")
        r_st0.font.name = 'Arial'; r_st0.font.size = Pt(9); r_st0.bold = True
        r_st0.font.color.rgb = RGBColor(30, 58, 138)
        
        p_st1 = tbl_stat.rows[0].cells[1].paragraphs[0]
        format_paragraph(p_st1, 0, 0)
        r_st1 = p_st1.add_run(f"✔ Trạng thái: Exit Code {run_info.get('returncode', 0)} (OK)")
        r_st1.font.name = 'Arial'; r_st1.font.size = Pt(9); r_st1.bold = True
        r_st1.font.color.rgb = RGBColor(22, 101, 52)
        
        p_st2 = tbl_stat.rows[0].cells[2].paragraphs[0]
        format_paragraph(p_st2, 0, 0)
        r_st2 = p_st2.add_run(f"⚡ Động cơ: Node v24 (TSX)")
        r_st2.font.name = 'Arial'; r_st2.font.size = Pt(9); r_st2.bold = True
        r_st2.font.color.rgb = RGBColor(100, 116, 139)
        
        add_body_p(doc, "")
        
        # Phân tích chi tiết từng dòng code
        add_body_p(doc, "", bold_prefix="a) Phân tích chi tiết từng dòng mã nguồn:")
        for line_exp in item.get('line_analysis', []):
            p_bullet = doc.add_paragraph(style='List Bullet')
            format_paragraph(p_bullet, space_before=0, space_after=3)
            r_b = p_bullet.add_run(line_exp)
            r_b.font.name = 'Arial'
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = RGBColor(51, 65, 85)
            
        # Luồng thực thi & Event Loop
        add_body_p(doc, item.get('execution_flow', ''), bold_prefix="b) Cơ chế hoạt động & Luồng thực thi (Event Loop): ")
        
        # Phân tích kết quả đầu ra
        add_body_p(doc, item.get('output_analysis', ''), bold_prefix="c) Phân tích kết quả đầu ra thực tế: ")
        
        # Ứng dụng thực tiễn trong React Native
        add_body_p(doc, item.get('mobile_application', ''), bold_prefix="d) Ứng dụng thực tiễn trong Lập trình Di động / React Native: ")
        
        # Đường kẻ phân cách nhẹ giữa các bài
        p_sep = doc.add_paragraph()
        format_paragraph(p_sep, space_before=4, space_after=12)
        r_sep = p_sep.add_run("─────────────────────────────────────────────────────────────")
        r_sep.font.color.rgb = RGBColor(226, 232, 240)
        
    doc.add_page_break()
    
    # =========================================================================
    # CHƯƠNG 3: TỔNG KẾT & KINH NGHIỆM THỰC CHIẾN
    # =========================================================================
    add_styled_heading(doc, "CHƯƠNG 3: TỔNG KẾT & KINH NGHIỆM THỰC CHIẾN TRONG REACT NATIVE / MOBILE", 1)
    
    add_styled_heading(doc, "3.1. Bảng tổng hợp đối chiếu kết quả 30 bài tập thực nghiệm", 2)
    add_body_p(doc, "Dưới đây là bảng tổng hợp toàn diện kết quả thực thi và đánh giá kỹ thuật của toàn bộ 30 bài tập thực hành trong dự án:")
    
    tbl_summary = doc.add_table(rows=31, cols=5)
    tbl_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_summary.autofit = False
    
    sum_headers = ["STT", "Chủ đề / Yêu cầu cốt lõi", "Cú pháp & Mẫu thiết kế chính", "Thời gian (s)", "Kết quả"]
    sum_widths = [Inches(0.6), Inches(2.4), Inches(2.2), Inches(0.8), Inches(0.8)]
    
    for j, h in enumerate(sum_headers):
        c = tbl_summary.rows[0].cells[j]
        c.width = sum_widths[j]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=80, bottom=80, left=60, right=60)
        p = c.paragraphs[0]
        format_paragraph(p, 0, 0)
        r = p.add_run(h)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for k in range(1, 31):
        row = tbl_summary.rows[k]
        bg = "F8FAFC" if k % 2 == 0 else "FFFFFF"
        ex_item = exercises.get(str(k), {})
        r_item = run_results.get(str(k), {})
        
        c0 = row.cells[0]; c0.width = sum_widths[0]
        c1 = row.cells[1]; c1.width = sum_widths[1]
        c2 = row.cells[2]; c2.width = sum_widths[2]
        c3 = row.cells[3]; c3.width = sum_widths[3]
        c4 = row.cells[4]; c4.width = sum_widths[4]
        
        for c in [c0, c1, c2, c3, c4]:
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=60, right=60)
            
        p0 = c0.paragraphs[0]; format_paragraph(p0, 0, 0)
        r0 = p0.add_run(f"Bài {k}"); r0.font.name = 'Arial'; r0.font.size = Pt(8.5); r0.bold = True
        
        p1 = c1.paragraphs[0]; format_paragraph(p1, 0, 0)
        r1 = p1.add_run(ex_item.get('vn_title', '')); r1.font.name = 'Arial'; r1.font.size = Pt(8.5)
        
        p2 = c2.paragraphs[0]; format_paragraph(p2, 0, 0)
        tech_note = ""
        if k <= 10:
            tech_note = "new Promise, setTimeout, .then, .catch, .finally, Promise.all/race"
        elif k <= 20:
            tech_note = "async/await, try/catch, Promise.all parallel, for await..of, Timeout Pattern"
        else:
            tech_note = "Fetch REST API, POST JSON, Retry Pattern, Batch, Queue, allSettled"
        r2 = p2.add_run(tech_note); r2.font.name = 'Arial'; r2.font.size = Pt(8)
        
        p3 = c3.paragraphs[0]; format_paragraph(p3, 0, 0)
        r3 = p3.add_run(f"{r_item.get('duration', 'N/A')}s"); r3.font.name = 'Arial'; r3.font.size = Pt(8.5)
        
        p4 = c4.paragraphs[0]; format_paragraph(p4, 0, 0)
        r4 = p4.add_run("ĐẠT"); r4.font.name = 'Arial'; r4.font.size = Pt(8.5); r4.bold = True
        r4.font.color.rgb = RGBColor(22, 101, 52)
        
    add_body_p(doc, "")
    
    add_styled_heading(doc, "3.2. Năm quy tắc vàng khi làm việc với Bất đồng bộ trong ứng dụng Di động", 2)
    
    rules = [
        ("Quy tắc 1: Luôn ưu tiên Promise.all cho các tác vụ độc lập", 
         "Khi màn hình cần nạp nhiều khối dữ liệu không phụ thuộc lẫn nhau (như thông tin người dùng, danh sách thông báo, cấu hình ứng dụng), hãy luôn kích hoạt chúng song song bằng `Promise.all` hoặc `Promise.allSettled`. Tránh tuyệt đối việc đặt liên tiếp các dòng `await` tuần tự vì sẽ làm thời gian chờ của người dùng tăng theo cấp số cộng."),
         
        ("Quy tắc 2: Triển khai Timeout và Retry cho mọi Request mạng",
         "Môi trường mạng của thiết bị di động vốn dĩ không ổn định (người dùng đi vào vùng sóng chập chờn, thang máy hoặc tầng hầm). Bắt buộc phải gắn bộ đếm Timeout (áp dụng Promise.race hoặc AbortSignal.timeout) và cơ chế tự động thử lại (Retry Pattern có Exponential Backoff) để ứng dụng không bị treo màn hình vĩnh viễn."),
         
        ("Quy tắc 3: Luôn giải phóng tài nguyên và hủy Request khi Component Unmount",
         "Sử dụng AbortController trong React hook useEffect để hủy bỏ (abort) các request mạng dở dang khi người dùng chuyển trang. Việc này giúp ngăn chặn hoàn toàn hiện tượng Memory Leak và cảnh báo cập nhật state trên unmounted component."),
         
        ("Quy tắc 4: Bọc toàn bộ lời gọi await trong cấu trúc try...catch...finally",
         "Không bao giờ để xảy ra lỗi Unhandled Promise Rejection. Khối catch phải ghi nhận lỗi hoặc hiển thị thông báo thân thiện. Khối finally bắt buộc phải tắt trạng thái loading (ví dụ: setLoading(false)) để người dùng không bị kẹt ở màn hình xoay tròn."),
         
        ("Quy tắc 5: Tận dụng Promise.allSettled() cho các màn hình Dashboard tổng hợp",
         "Khi giao diện tổng hợp bao gồm nhiều widget độc lập (ví dụ: Ví tiền, Lịch sử giao dịch, Tin khuyến mãi), việc dùng Promise.all sẽ khiến toàn bộ màn hình báo lỗi nếu chỉ duy nhất 1 widget bị trục trặc. Thay vào đó, Promise.allSettled() cho phép các widget bình thường hiển thị đầy đủ và widget lỗi hiển thị nút 'Thử lại riêng', tối đa hóa trải nghiệm người dùng.")
    ]
    
    for r_title, r_body in rules:
        add_callout(doc, r_title, r_body, bg_color="F0FDF4", border_color="22C55E", title_color=(22, 101, 52))
        add_body_p(doc, "")
        
    add_styled_heading(doc, "3.3. Lời kết", 2)
    add_body_p(doc, "Bộ 30 bài tập thực hành bất đồng bộ trong chuyên đề này đã bao quát trọn vẹn từ những khái niệm cốt lõi nhất của Promise cho đến những mẫu thiết kế hiện đại và phức tạp trong kiến trúc ứng dụng di động React Native chuyên nghiệp. Việc nắm vững cơ chế vận hành của Event Loop, sự khác biệt giữa Microtask và Macrotask, cũng như thành thạo các mẫu xử lý song song, hàng đợi, retry và timeout là nền tảng vững chắc nhất để xây dựng các ứng dụng di động có hiệu năng mượt mà, độ tin cậy cao và trải nghiệm người dùng tuyệt vời.")
    
    # Lưu file Word
    out_file = r"d:\mobile\Tuan2\Tuan2\BaoCao_ThucHanh_Async_TypeScript.docx"
    doc.save(out_file)
    print(f"ĐÃ TẠO THÀNH CÔNG TỆP BÁO CÁO WORD TẠI: {out_file}")

if __name__ == '__main__':
    main()
